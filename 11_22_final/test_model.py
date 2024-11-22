
import openpyxl
from KKT_Module.ksoc_global import kgl
from KKT_Module.Configs import SettingConfigs
from KKT_Module.SettingProcess.SettingProccess import SettingProc, ConnectDevice, ResetDevice
from KKT_Module.DataReceive.DataReciever import RawDataReceiver, HWResultReceiver, FeatureMapReceiver,MultiResult4168BReceiver
import time
import cv2
import numpy as np
import tensorflow as tf
from tensorflow.keras.models import Sequential, load_model
from tensorflow.keras.layers import Conv2D, MaxPooling2D, Flatten, Dense
from sklearn.model_selection import train_test_split
import serial
import datetime

# 設定串行通信參數
arduino = serial.Serial(port='COM6', baudrate=9600, timeout=.1)

def send_person_count(count):
    if count<0: count = 0
    byte_str = bytes(str(count), 'utf-8') + b'\n'
    arduino.write(byte_str)
    # time.sleep(1.05)

def add_to_dict(dic, val):
    if val not in dic:
        dic[val] = 0
    else:
        dic[val] += 1
def dict_max(dic: dict) -> int:
    max_value = max(dic.values())
    max_keys = [key for key, value in dic.items() if value == max_value]
    return max(max_keys)


def connect():
    connect = ConnectDevice()
    connect.startUp()                       # Connect to the device
    reset = ResetDevice()
    reset.startUp()                         # Reset hardware register

def startSetting():
    SettingConfigs.setScriptDir("K60168-Test-00256-008-v0.0.8-20230717_480cm")  # Set the setting folder name
    ksp = SettingProc()                 # Object for setting process to setup the Hardware AI and RF before receive data
    ksp.startUp(SettingConfigs)             # Start the setting process
    # ksp.startSetting(SettingConfigs)        # Start the setting process in sub_thread

def startLoop():
    # kgl.ksoclib.switchLogMode(True)
    # R = RawDataReceiver(chirps=32)
    # wr = openpyxl.Workbook()
    # frameCount = 500
    whileCount = 0
    mean_count = 0
    mean_mean = 0
    # mean_mean2 = 0
    # cv2.namedWindow("draw",cv2.WINDOW_NORMAL)
    # cv2.namedWindow("number",cv2.WINDOW_NORMAL)
    cv2.namedWindow("mean",cv2.WINDOW_NORMAL)
    # cv2.namedWindow("move_list",cv2.WINDOW_NORMAL)
    # cv2.namedWindow("draw2",cv2.WINDOW_NORMAL)
    n = np.zeros((32, 32, 3), np.uint8)
    model = load_model('my_model.h5')
    # model2 = load_model('my_model.h5')
    # model.load_weights('./b16_lr001.h5')
    # model2.load_weights('./b16_lr001_2000.h5')
    model.load_weights('./weight_rditrain125.h5')

    # s = wr.create_sheet("raw_data")
    # s = wr.create_sheet('rdi_data')
    # Receiver for getting Raw data
    # R = MultiResult4168BReceiver()
    R = FeatureMapReceiver(chirps=32)       # Receiver for getting RDI PHD map
    # R = HWResultReceiver()                  # Receiver for getting hardware results (gestures, Axes, exponential)
    # buffer = DataBuffer(100)                # Buffer for saving latest frames of data
    R.trigger(chirps=32)                             # Trigger receiver before getting the data
    time.sleep(0.5)
    print('# ======== Start getting gesture ===========')
    # dic = dict()
    # move_list = []
    # move_list_temp = 0
    # previousMillis = 0
    # interval = 3
    # personCount = 0
    # lastPersonCount = 0
    while True:
        
        # now = datetime.datetime.now()
        # currentMillis = now.second
        
    # while whileCount<frameCount:                             # loop for getting the data
        res = R.getResults()                # Get data from receiver
        if res is None:
            continue
        whileCount+=1
        # print(whileCount)
        #draw
        # for i in range(1, 33):
        #     for j in range(1, 33):
        #         value = res[0][j-1][i-1]
        #         hls = (value*360/1200, 50, 85)
        #         hls_conv = (hls[0]/2, hls[1]*2.55, hls[2]*2.55)
        #         rgb = cv2.cvtColor(np.uint8([[hls_conv]]), cv2.COLOR_HLS2RGB)[0][0]
        #         n[j-1, i-1] = rgb
        a = list(res[0])
        new_frame = a/np.max(a)
        new_frame = np.array(new_frame)
        new_frame = new_frame.reshape(1, 32, 32, 1)  # Reshape for model input
        
        # pridicted_objects2 = model2.predict(new_frame)
        predicted_objects = model.predict(new_frame)
        print(f'Predicted number of objects: {predicted_objects[0][0]}')
        
        # text_num = np.zeros((100,100,3),np.uint8)
        # cv2.putText(text_num,f"{round(predicted_objects[0][0])}",(20,70),cv2.FONT_HERSHEY_SIMPLEX,3,(255,255,255),12)
        # cv2.imshow('number',text_num)
        # if whileCount>=51:
        #     time.sleep(1/2)
        #     send_data = sum(move_list)//50
        #     # if move_list_temp != send_data:
        #     if True:
        #         mean_draw2 = np.zeros((100,100,3),np.uint8)
        #         cv2.putText(mean_draw2,f"{send_data}",[20,70],cv2.FONT_HERSHEY_SIMPLEX,3,(255,255,255),12)
        #         cv2.imshow("move_list",mean_draw2)
        #         if (send_data >= lastPersonCount or currentMillis - previousMillis >= interval) :
        #             previousMillis = currentMillis
        #             lastPersonCount = send_data
        #         else:
        #             send_data = lastPersonCount
        #         send_person_count(send_data)
        #     # move_list_temp = send_data
        #     del move_list[0]
        # move_list.append(round(predicted_objects[0][0]))
        
        # mean_mean2 += round(pridicted_objects2[0][0])

        mean_count += 1
        # add_to_dict(dic, round(predicted_objects[0][0]))
        mean_mean += round(predicted_objects[0][0])
        frame_unit = 25
        if mean_count == frame_unit:
            # dic_max = dict_max(dic)
            # dic = dict()
            mean_count = 0
            mean_draw = np.zeros((100,100,3),np.uint8)
            final_people = mean_mean // frame_unit
            cv2.putText(mean_draw,f"{final_people}",[20,70],cv2.FONT_HERSHEY_SIMPLEX,3,(255,255,255),12)
            print("detect people: ", final_people)
            send_person_count(final_people)
            mean_mean = 0
            cv2.imshow("mean",mean_draw)

            # mean_draw2 = np.zeros((100,100,3),np.uint8)
            # # final_people2 = mean_mean2 // frame_unit
            # cv2.putText(mean_draw2,f"{sum(move_list)//25}",[20,70],cv2.FONT_HERSHEY_SIMPLEX,3,(255,255,255),12)
            # # print("detect people2: ", final_people2)
            # # send_person_count(final_people*2+2)
            # mean_mean2 = 0
            # cv2.imshow("b16_lr001",mean_draw2)
        
        # cv2.imshow('draw2',n2)
        # cv2.imshow("draw", n)
        key = cv2.waitKey(1)
        if key == ord(' '):
            break

    #     #RDI
        # for i in res[0]:
            # s.append(list(i))

    #     # time.sleep(0.05)
    #     '''
    #     Application for the data.
    #     '''
    cv2.destroyAllWindows()
    # wr.save("rdi5test.xlsx")
    print('\a')
def main():
    kgl.setLib()

    # kgl.ksoclib.switchLogMode(True)

    connect()                               # First you have to connect to the device

    startSetting()                         # Second you have to set the setting configs

    startLoop()                             # Last you can continue to get the data in the loop

if __name__ == '__main__':
    main()
